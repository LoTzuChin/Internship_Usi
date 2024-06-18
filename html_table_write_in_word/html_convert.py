from bs4 import BeautifulSoup
from docx import Document

def replace_chars(text):
    """
    將 input的字串進行處理，取代掉不需要的字符
    """
    return text.replace(',', '').replace('[', '').replace(']', '').replace("'", "")

def write_to_html(text):
    """
    將處理過的字串寫進 .html中
    """
    with open('output.html', 'w') as f:
        f.write(text)

def get_max_cols(table):
    """
    取得行數
    """
    max_cols = 0
    for row in table.find_all("tr"):
        cols = len(row.find_all(["th", "td"]))
        if cols > max_cols:
            max_cols = cols
    return max_cols


def get_max_rows(table):
    """
    取得列數
    """
    return len(table.find_all("tr"))


def process_table(doc, table):
    """
    將 html內容表格轉入 word中
    """
    max_cols = get_max_cols(table)
    max_rows = get_max_rows(table)

    # 獲取所有儲存格內容
    cells = [cell for row in table.find_all("tr") for cell in row.find_all(["th", "td"])]

    # 創建一個新表格
    doc_table = doc.add_table(rows=max_rows, cols=max_cols)

    for cell in cells:
        # 獲取表格內容
        text = cell.get_text(strip=True)
        # 如果儲存格內容為空，則將其替換為一個空白鍵
        if not text:
            text = ' '
        # 獲得合併儲存格狀況
        rowspan = int(cell.get('rowspan', 1))
        colspan = int(cell.get('colspan', 1))

        # 將表格內容寫到對應位置
        for row_idx in range(max_rows):
            for col_idx in range(max_cols):
                if not doc_table.cell(row_idx, col_idx).text:
                    for r in range(row_idx, row_idx + rowspan):
                        for c in range(col_idx, col_idx + colspan):
                            if r != row_idx or c != col_idx:
                                doc_table.cell(r, c).merge(doc_table.cell(row_idx, col_idx))
                    doc_table.cell(row_idx, col_idx).text = text
                    break
            else:
                continue
            break



if __name__ == '__main__':

    # 處理輸入進來的列表
    input_text = input("請輸入一段文字: ")
    formatted_text = replace_chars(input_text)
    write_to_html(formatted_text)

    # 將表格寫進 word
    # 讀 HTML文件並解析為 BeautifulSoup物件
    with open("output.html", "r", encoding="utf-8") as f:
        html_content = f.read()
    soup = BeautifulSoup(html_content, "html.parser")

    # 創建新的 word
    doc = Document()

    # 找到 HTML 中的表格
    table = soup.find("table")

    # 處理表格，寫進 word
    if table:
        process_table(doc, table)
    # 保存 word檔
    doc.save("output.docx")
